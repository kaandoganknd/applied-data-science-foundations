"""Validate every W9-W10 answer-key statement against the unchanged reference data."""

import argparse
import calendar
import csv
import getpass
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import pymysql
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parent
SOURCE = ROOT / 'sources/MySQL Activities W9-W10 - Solutions.docx'
SQL = ROOT / 'sql/answer_key_part2.sql'
DATA = REPO / 'week-08-mysql-creditcard/data/creditcard_insert.csv'
SCHEMA = REPO / 'week-08-mysql-creditcard/evidence/show-create-table.sql'
FIELDS = ['CreditcardNum', 'Creditcard_company', 'Creditcard_type', 'Credit_Limit',
          'Totalspent', 'City', 'CardHolder', 'Issue_Date']


def plain(value):
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, date):
        return value.isoformat()
    return value


def bag(rows):
    return Counter(tuple(row) for row in rows)


def table(columns, rows):
    def cell(value):
        return str(plain(value) if value is not None else 'NULL').replace('|', '\\|')
    return '\n'.join(['| ' + ' | '.join(columns) + ' |',
                      '| ' + ' | '.join(['---'] * len(columns)) + ' |'] +
                     ['| ' + ' | '.join(map(cell, row)) + ' |' for row in rows])


def plus_months(day, months):
    year, month = divmod(day.year * 12 + day.month - 1 + months, 12)
    return date(year, month + 1, min(day.day, calendar.monthrange(year, month + 1)[1]))


def load_questions():
    questions, active = [], False
    for paragraph in Document(SOURCE).paragraphs:
        text = paragraph.text.strip()
        if text == 'PART 2':
            active = True
        if not active:
            continue
        prop = paragraph._p.pPr
        num = prop.numPr if prop is not None else None
        if num is not None and int(num.ilvl.val) == 0:
            questions.append(dict(number=len(questions) + 1, text=text, statements=[]))
        elif text.startswith(('SELECT ', 'SET ')):
            questions[-1]['statements'].append(text)
    assert len(questions) == 26
    statements = [sql for q in questions for sql in q['statements']]
    assert sum(sql.startswith('SELECT ') for sql in statements) == 29
    code = re.sub(r'^--[^\n]*', '', SQL.read_text(), flags=re.M).strip()
    assert code.startswith('USE financial_db;')
    actual = [piece.strip() + ';' for piece in code.split(';')[1:] if piece.strip()]
    assert actual == statements, 'Runnable file differs from supplied answer-key statements'
    old_source = REPO / 'week-09-mysql-creditcard-queries/sources/MySQL Activities W8- Solutions.docx'
    earlier = [p.text.strip() for p in Document(old_source).paragraphs
               if p.text.strip().startswith('SELECT ') and p.text.strip().endswith(';')]
    assert earlier == [sql for q in questions[:13] for sql in q['statements']]
    return questions, old_source


def load_records():
    with DATA.open(newline='') as stream:
        rows = list(csv.DictReader(stream))
    for row in rows:
        row['CreditcardNum'] = int(row['CreditcardNum'])
        row['Credit_Limit'] = Decimal(row['Credit_Limit'])
        row['Totalspent'] = Decimal(row['Totalspent'])
        row['Issue_Date'] = date.fromisoformat(row['Issue_Date'])
    return sorted(rows, key=lambda r: r['CreditcardNum'])


def expected_results(rows):
    def project(records, *fields):
        return [tuple(r[f] for f in fields) for r in records]
    city, holder, year, month = (defaultdict(list) for _ in range(4))
    for row in rows:
        city[row['City']].append(row)
        holder[row['CardHolder']].append(row)
        year[row['Issue_Date'].year].append(row)
        month[row['Issue_Date'].month].append(row)
    remaining = lambda r: r['Credit_Limit'] - r['Totalspent']
    max_limit = max(r['Credit_Limit'] for r in rows)
    average = (sum(r['Credit_Limit'] for r in rows)/len(rows)).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
    return {
        1: [(len(rows),)],
        2: [(len(set(r['Creditcard_company'] for r in rows)),)],
        3: [(x,) for x in set(r['Creditcard_company'] for r in rows)],
        4: project(rows, 'CreditcardNum', 'Credit_Limit'),
        5: project([r for r in rows if r['Credit_Limit'] == max_limit], 'CreditcardNum', 'Credit_Limit'),
        6: project([r for r in rows if r['City'] == 'London'], 'CreditcardNum', 'City'),
        7: project([r for r in rows if r['City'].startswith('Aber') and r['Credit_Limit'] > 5000], 'CreditcardNum', 'City', 'Credit_Limit'),
        8: project([r for r in rows if r['Credit_Limit'] in [1200,3000,4500]], 'CreditcardNum', 'Credit_Limit'),
        9: project([r for r in rows if r['City'] != 'London'], 'CreditcardNum', 'City'),
        10: [('Mastercard', sum(r['Creditcard_type'] == 'Mastercard' for r in rows))],
        11: [(key, len(value)) for key,value in city.items()],
        12: [(key, len(value)) for key,value in holder.items()],
        13: [(key,min(r['Credit_Limit'] for r in value),max(r['Credit_Limit'] for r in value)) for key,value in city.items()],
        14: [(key,min(r['Credit_Limit'] for r in value),max(r['Credit_Limit'] for r in value),sum(r['Totalspent'] for r in value)) for key,value in city.items()],
        15: [(key,min(r['Totalspent'] for r in value),max(r['Totalspent'] for r in value)) for key,value in holder.items()],
        16: [(key,sum(r['Totalspent'] for r in value)) for key,value in holder.items()],
        17: [(r['CreditcardNum'],remaining(r)) for r in rows],
        18: [(key,len(value)) for key,value in year.items()],
        19: [(key,len(value)) for key,value in month.items()],
        20: project([r for r in rows if r['Issue_Date'] > date(2020,1,1)], 'CreditcardNum', 'Issue_Date'),
        21: [(key,sum(remaining(r) for r in value)) for key,value in holder.items()],
        22: [(key,sum(remaining(r) for r in value)) for key,value in year.items()],
        23: [(key,sum(remaining(r) for r in value)) for key,value in month.items()],
        24: [(r['CreditcardNum'],r['Issue_Date'],plus_months(r['Issue_Date'],18)) for r in rows],
        25: [(r['CreditcardNum'],r['Issue_Date'],plus_months(r['Issue_Date'],18),plus_months(r['Issue_Date'],17)) for r in rows],
        26: [(r['CardHolder'],r['Credit_Limit'],average) for r in rows if r['Credit_Limit'] > average],
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--host', default='127.0.0.1')
    parser.add_argument('--port', type=int, default=3306)
    parser.add_argument('--socket')
    parser.add_argument('--user', default='root')
    parser.add_argument('--ask-password', action='store_true')
    args = parser.parse_args()
    questions, old_source = load_questions()
    records = load_records()
    assert len(records) == 15
    expected = expected_results(records)
    results, edge_cases = [], []
    connection = pymysql.connect(host=args.host, port=args.port, unix_socket=args.socket,
                                 user=args.user, database='financial_db', charset='utf8mb4',
                                 password=getpass.getpass('MySQL password: ') if args.ask_password else '',
                                 autocommit=True)
    try:
        with connection.cursor() as cursor:
            cursor.execute('SET SESSION TRANSACTION READ ONLY')
            connection.begin()
            cursor.execute('SELECT VERSION(), @@sql_mode, @@lower_case_table_names')
            version, mode, case_mode = cursor.fetchone()
            cursor.execute('SHOW CREATE TABLE creditcard')
            ddl = cursor.fetchone()[1]
            assert ddl + ';\n' == SCHEMA.read_text()
            cursor.execute('SELECT * FROM creditcard ORDER BY CreditcardNum')
            before = cursor.fetchall()
            assert [c[0] for c in cursor.description] == FIELDS
            assert bag(before) == bag(tuple(r[f] for f in FIELDS) for r in records)
            for question in questions:
                number, variant = question['number'], 0
                for sql in question['statements']:
                    assert sql.count(';') == 1 and sql.endswith(';')
                    if sql.startswith('SET '):
                        assert sql.startswith('SET @AvgCredLimit = ROUND(')
                        cursor.execute(sql)
                        continue
                    assert sql.startswith('SELECT ')
                    cursor.execute(sql)
                    columns = [c[0] for c in cursor.description]
                    actual = cursor.fetchall()
                    cursor.execute('SHOW WARNINGS')
                    assert not cursor.fetchall(), (number, 'Unexpected warning')
                    want = expected[number]
                    if number == 10 and 'Visa' in sql:
                        want = [(t, sum(r['Creditcard_type'] == t for r in records)) for t in ['Mastercard','Visa']]
                    assert bag(actual) == bag(want), (number, actual, want)
                    if number == 4:
                        assert [r[1] for r in actual] == sorted([r['Credit_Limit'] for r in records], reverse=True)
                    variant += 1
                    results.append(dict(question=number, variant=variant, title=question['text'],
                                        sql=sql, columns=columns, rows=actual, row_count=len(actual), check='PASS'))
            for value in ['2020-02-29','2020-08-31','2020-01-01']:
                cursor.execute('SELECT DATE_ADD(%s, INTERVAL 18 MONTH), DATE_ADD(%s, INTERVAL 17 MONTH)', (value,value))
                expiry, replacement = cursor.fetchone()
                actual = (str(expiry),str(replacement))
                want = tuple(plus_months(date.fromisoformat(value), n).isoformat() for n in [18,17])
                assert actual == want
                edge_cases.append(dict(issue_date=value, expiry=actual[0], replacement=actual[1], check='PASS'))
            cursor.execute("SELECT DATE('2020-01-01') > '2020-01-01', DATE('2020-01-02') > '2020-01-01'")
            assert cursor.fetchone() == (0,1)
            cursor.execute('SELECT * FROM creditcard ORDER BY CreditcardNum')
            assert cursor.fetchall() == before
            cursor.execute('SHOW CREATE TABLE creditcard')
            assert cursor.fetchone()[1] == ddl
    finally:
        connection.rollback()
        connection.close()
    total_remaining = sum(r['Credit_Limit'] - r['Totalspent'] for r in records)
    assert sum(r[1] for r in expected[21]) == sum(r[1] for r in expected[22]) == sum(r[1] for r in expected[23]) == total_remaining
    assert sum(r[3] for r in expected[14]) == sum(r['Totalspent'] for r in records)
    assert len(results) == 29

    out = ROOT / 'evidence/answer-key'
    out.mkdir(exist_ok=True)
    report = ['# W9-W10 answer key: complete executed results', '',
              f'MySQL {version}; all 26 Part 2 questions and 29 SELECT variants. Every result cell was checked against independent calculations on the 15 reference records.', '',
              'AK numbering follows the W9-W10 solutions document. SQL statements are unchanged; row order is only guaranteed where the source specifies ORDER BY.', '',
              'The source setup is in Week 8. Differences from the earlier question sheet are explained in [coverage notes](../../answer-key-coverage.md).', '']
    previous = None
    for result in results:
        number = result['question']
        if number != previous:
            report += [f'## AK{number:02d}. {result["title"]}', '']
            if number == 26:
                report += ['```sql', questions[25]['statements'][0], '```', '']
        if number in [5,10]:
            report += [f'### Variant {result["variant"]}', '']
        report += ['```sql',result['sql'],'```','',table(result['columns'],result['rows']),'']
        filename = f'q{number:02d}-v{result["variant"]:02d}.csv'
        result['csv'] = filename
        with (out/filename).open('w',newline='') as stream:
            writer = csv.writer(stream,lineterminator='\n')
            writer.writerow(result['columns'])
            writer.writerows(tuple(plain(v) for v in row) for row in result['rows'])
        previous = number
    report += ['## Date checks', '', table(['Issue date','18-month expiry','17-month replacement'],
               [(r['issue_date'],r['expiry'],r['replacement']) for r in edge_cases]), '',
               'The strict 2020-01-01 cutoff excludes that date and includes 2020-01-02. Month-end dates can make issue date + 17 months differ from expiry date - 1 month; AK25 preserves the supplied formula.', '']
    (out/'query-results.md').write_text('\n'.join(report))
    inputs = [SOURCE,SQL,DATA,SCHEMA,old_source,Path(__file__).resolve()]
    manifest = dict(verified_at_utc=datetime.now(timezone.utc).isoformat(), mysql_version=version,
                    sql_mode=mode, lower_case_table_names=case_mode, source_rows=15,
                    part2_questions_checked=26, select_variants_checked=29,
                    statement_text_matches_source=True, first_13_questions_match_previous_key=True,
                    independent_result_checks_passed=True, schema_and_records_unchanged=True,
                    transaction_read_only=True, date_add_cases=edge_cases, strict_cutoff_check='PASS',
                    remaining_credit_reconciles_by_holder_year_and_month=True,
                    input_sha256={str(p.relative_to(REPO)):hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},
                    results=results)
    (out/'validation.json').write_text(json.dumps(manifest,indent=2,default=plain,ensure_ascii=False)+'\n')
    print(f'MySQL {version}: all 26 questions / 29 SELECT variants passed; source SQL preserved; dates and totals verified; schema and 15 records unchanged.')


if __name__ == '__main__':
    main()
