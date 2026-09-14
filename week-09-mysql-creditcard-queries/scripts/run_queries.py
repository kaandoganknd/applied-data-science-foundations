"""Run all thirteen Part 2 questions and every supplied answer-key alternative."""

import argparse
import csv
import getpass
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pymysql
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SETUP = ROOT.parent / 'week-08-mysql-creditcard'
FIELDS = ['CreditcardNum', 'Creditcard_company', 'Creditcard_type', 'Credit_Limit',
          'Totalspent', 'City', 'CardHolder', 'Issue_Date']


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def plain(value):
    if isinstance(value, Decimal):
        return format(value, '.2f')
    if isinstance(value, date):
        return value.isoformat()
    return value


def table(columns, rows):
    def cell(value):
        return str(plain(value) if value is not None else 'NULL').replace('|', '\\|')
    return '\n'.join(['| ' + ' | '.join(columns) + ' |',
                      '| ' + ' | '.join(['---'] * len(columns)) + ' |'] +
                     ['| ' + ' | '.join(map(cell, row)) + ' |' for row in rows])


def load_questions():
    doc = Document(ROOT / 'sources/MySQL Activities W8-W9.docx')
    questions = []
    for paragraph in doc.paragraphs:
        text = ' '.join(paragraph.text.split())
        if not text:
            continue
        prop = paragraph._p.pPr
        numbering = prop.numPr if prop is not None else None
        if numbering is not None and int(numbering.ilvl.val) == 0:
            questions.append({'number': len(questions) + 1, 'text': text, 'subtasks': []})
        elif numbering is not None:
            questions[-1]['subtasks'].append(text)
        elif questions:
            questions[-1]['text'] += ' ' + text
    assert len(questions) == 26 and len(questions[2]['subtasks']) == 6
    assert questions[12]['text'] == 'List all credit cards that were not issued in London?'
    return questions[:17]


def answer_key():
    questions = []
    active = False
    for p in Document(ROOT / 'sources/MySQL Activities W8- Solutions.docx').paragraphs:
        text = p.text.strip()
        if text == 'PART 2':
            active = True
        if not active:
            continue
        prop = p._p.pPr
        numbering = prop.numPr if prop is not None else None
        if numbering is not None and int(numbering.ilvl.val) == 0:
            questions.append({'number': len(questions) + 1, 'text': text, 'sql': []})
        elif text.startswith('SELECT '):
            questions[-1]['sql'].append(text)
    assert len(questions) == 13 and sum(len(q['sql']) for q in questions) == 16
    return questions


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--host', default='127.0.0.1')
    parser.add_argument('--port', type=int, default=3306)
    parser.add_argument('--socket')
    parser.add_argument('--user', default='root')
    parser.add_argument('--ask-password', action='store_true')
    args = parser.parse_args()
    source = SETUP / 'data/creditcard_insert.csv'
    with source.open(newline='') as stream:
        expected = list(csv.DictReader(stream))
    for row in expected:
        row['CreditcardNum'] = int(row['CreditcardNum'])
        for field in ['Credit_Limit', 'Totalspent']:
            row[field] = Decimal(row[field])
    expected.sort(key=lambda row: row['CreditcardNum'])
    assert len(expected) == 15
    sql_files = [ROOT / 'sql' / name for name in ['questions_05_to_13.sql', 'questions_14_to_17.sql']]
    blocks = re.findall(r'-- Q(\d{2}): ([^\n]+)\n(.*?;)', '\n'.join(p.read_text() for p in sql_files), re.S)
    assert [int(n) for n, _, _ in blocks] == list(range(5, 18))
    assert all(query.lstrip().startswith('SELECT ') for _, _, query in blocks)
    questions = load_questions()
    companies = sorted({row['Creditcard_company'] for row in expected})
    highest = max(row['Credit_Limit'] for row in expected)
    card_types = Counter(row['Creditcard_type'] for row in expected)
    cities = Counter(row['City'] for row in expected)
    holders = Counter(row['CardHolder'] for row in expected)
    limits = defaultdict(list)
    for row in expected:
        limits[row['City']].append(row['Credit_Limit'])
    to_rows = lambda records: [tuple(plain(row[field]) for field in FIELDS) for row in records]
    expected_results = {
        5: [(len(expected),)],
        6: [(len(companies),)],
        7: [(name,) for name in companies],
        8: to_rows(sorted(expected, key=lambda row: (-row['Credit_Limit'], row['CreditcardNum']))),
        9: to_rows([row for row in expected if row['Credit_Limit'] == highest]),
        10: to_rows([row for row in expected if row['City'] == 'London']),
        11: to_rows([row for row in expected if row['City'] == 'Aberdeen' and row['Credit_Limit'] > 5000]),
        12: to_rows([row for row in expected if row['Credit_Limit'] in [1200, 3000, 4500]]),
        13: to_rows([row for row in expected if row['City'] != 'London']),
        14: [('Mastercard', card_types['Mastercard'])],
        15: sorted(cities.items()),
        16: sorted(holders.items()),
        17: [(city, plain(min(values)), plain(max(values))) for city, values in sorted(limits.items())],
    }
    conn = pymysql.connect(host=args.host, port=args.port, unix_socket=args.socket,
                           user=args.user, database='financial_db', charset='utf8mb4',
                           password=getpass.getpass('MySQL password: ') if args.ask_password else '',
                           autocommit=True)
    results = []
    reference_results = []
    references = answer_key()
    try:
        with conn.cursor() as cursor:
            cursor.execute('SET SESSION TRANSACTION READ ONLY')
            conn.begin()
            cursor.execute('SELECT VERSION(), @@lower_case_table_names, @@sql_mode')
            version, case_mode, sql_mode = cursor.fetchone()
            cursor.execute('SHOW CREATE TABLE creditcard')
            ddl_before = cursor.fetchone()[1]
            assert ddl_before + ';\n' == (SETUP / 'evidence/show-create-table.sql').read_text(), 'Schema differs from the verified setup'
            cursor.execute('DESCRIBE creditcard')
            schema_columns = [column[0] for column in cursor.description]
            schema_rows = cursor.fetchall()
            cursor.execute('SELECT * FROM creditcard ORDER BY CreditcardNum')
            before = cursor.fetchall()
            assert [column[0] for column in cursor.description] == FIELDS
            assert [tuple(map(plain, row)) for row in before] == to_rows(expected), 'Table differs from the 15 source records; no data was changed'
            for number, title, query in blocks:
                cursor.execute(query)
                columns = [column[0] for column in cursor.description]
                rows = [tuple(map(plain, row)) for row in cursor.fetchall()]
                cursor.execute('SHOW WARNINGS')
                assert not cursor.fetchall(), f'Q{number}: unexpected warning'
                assert rows == expected_results[int(number)], f'Q{number}: result mismatch'
                results.append(dict(question=int(number), title=title, sql=query,
                                    columns=columns, rows=rows, row_count=len(rows), check='PASS'))
            for question in references:
                own = next(r for r in results if r['question'] == question['number'] + 4)
                for variant, sql in enumerate(question['sql'], 1):
                    cursor.execute(sql)
                    columns = [column[0] for column in cursor.description]
                    rows = [tuple(map(plain, row)) for row in cursor.fetchall()]
                    cursor.execute('SHOW WARNINGS')
                    assert not cursor.fetchall(), 'Unexpected warning in answer key'
                    if question['number'] == 10 and variant > 1:
                        projected = [(kind, card_types[kind]) for kind in ['Mastercard', 'Visa']]
                    else:
                        projected = [tuple(row[own['columns'].index(field)] if field in own['columns'] else row[i]
                                           for i, field in enumerate(columns)) for row in own['rows']]
                    assert Counter(rows) == Counter(projected), (question['number'], variant, 'Answer-key mismatch')
                    if question['number'] == 4:
                        returned_limits = [Decimal(r[1]) for r in rows]
                        assert returned_limits == sorted(returned_limits, reverse=True)
                    reference_results.append(dict(part2_question=question['number'], variant=variant,
                                                  sql=sql, columns=columns, rows=rows, check='PASS'))
            cursor.execute('SELECT * FROM creditcard ORDER BY CreditcardNum')
            assert cursor.fetchall() == before, 'Records changed'
            cursor.execute('SHOW CREATE TABLE creditcard')
            assert cursor.fetchone()[1] == ddl_before, 'Schema changed'
    finally:
        conn.rollback()
        conn.close()

    out = ROOT / 'evidence'
    out.mkdir(exist_ok=True)
    docs = ROOT / 'docs'
    docs.mkdir(exist_ok=True)
    question_text = ['# Source questions and answer-key mapping', '',
                     'Original questions 1–4 are setup. Original Q5–17 correspond to Part 2 Q1–13 in the answer key; the six subparts of Q3 are retained. Original Q18–26 continue in Week 10.', '']
    for item in questions:
        question_text += [f"## {item['number']}. {item['text']}", '']
        question_text += [f'- {part}' for part in item['subtasks']]
        if item['subtasks']:
            question_text.append('')
    (docs / 'questions.md').write_text('\n'.join(question_text).rstrip() + '\n')
    report = ['# Executed Creditcard queries', '',
              f'MySQL {version}. Read-only queries on the 15 unchanged reference records in `financial_db.creditcard`.', '',
              'Questions 1–4 use the existing setup; the structure below was read again in this run. Original Q5–17 cover all thirteen Part 2 questions, in source order with complete outputs.', '',
              '## Table structure (Q2–Q3)', '', '```sql', 'DESCRIBE creditcard;', '```', '',
              table(schema_columns, schema_rows), '']
    for result in results:
        report += [f"## Q{result['question']:02d}. {result['title']}", '',
                   '```sql', result['sql'], '```', '',
                   table(result['columns'], result['rows']), '',
                   f"Rows returned: {result['row_count']}.", '']
        if result['question'] == 8:
            report += ['The answer key confirms descending sorting for Q8; Q9 separately selects the maximum-limit card(s).', '']
        with (out / f"q{result['question']:02d}.csv").open('w', newline='') as stream:
            writer = csv.writer(stream, lineterminator='\n')
            writer.writerow(result['columns'])
            writer.writerows(result['rows'])
    (out / 'query-results.md').write_text('\n'.join(report))
    reference_report = ['# Answer-key queries and complete results', '',
                        'Every SELECT statement is extracted unchanged from the supplied Word answer key, including all alternatives. Part 2 numbering restarts at one.', '']
    reference_sql = ['USE financial_db;', '']
    for result in reference_results:
        label = f"Part 2 Q{result['part2_question']:02d}, variant {result['variant']}"
        reference_report += ['## ' + label, '', '```sql', result['sql'], '```', '', table(result['columns'], result['rows']), '']
        reference_sql += ['-- ' + label, result['sql'], '']
    (out / 'answer-key-results.md').write_text('\n'.join(reference_report))
    reference_path = ROOT / 'sql/answer_key_part2.sql'
    reference_path.write_text('\n'.join(reference_sql).rstrip() + '\n')
    (out / 'table-structure.sql').write_text(ddl_before + ';\n')
    manifest = dict(verified_at_utc=datetime.now(timezone.utc).isoformat(),
                    mysql_version=version, lower_case_table_names=case_mode, sql_mode=sql_mode,
                    schema='financial_db', table='creditcard', reference_rows=15,
                    source_question_numbers=list(range(1, 18)), newly_executed_questions=list(range(5, 18)),
                    answer_key_part2_questions=list(range(1, 14)), answer_key_variants_checked=len(reference_results),
                    passed_query_checks=len(results), failed_query_checks=0,
                    comparison='Every result cell compared with independent Python calculations on the reference INSERT CSV',
                    schema_matches_reference=True, records_match_reference=True,
                    schema_and_records_unchanged=True, transaction_read_only=True,
                    input_sha256={str(path.relative_to(ROOT.parent)): digest(path) for path in [
                        ROOT / 'sources/MySQL Activities W8-W9.docx',
                        ROOT / 'sources/MySQL Activities W8- Solutions.docx', *sql_files, reference_path, source,
                        SETUP / 'sql/full_solution.sql', SETUP / 'evidence/show-create-table.sql']},
                    results=results, answer_key_results=reference_results)
    (out / 'validation.json').write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + '\n')
    print(f'MySQL {version}: 13/13 query checks and 16 answer-key variants passed; schema and 15 source records unchanged.')
    for result in results:
        print(f"Q{result['question']:02d}: {result['row_count']} output rows")


if __name__ == '__main__':
    main()
